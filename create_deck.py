import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck(filename="ExamAlert_AI_HACKDAY_1.0.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette - Modern Tech Dark Mode
    BG_COLOR = RGBColor(11, 15, 25)       # #0B0F19 Slate Darkest
    CARD_BG = RGBColor(23, 32, 51)        # #172033 Deep Card
    CARD_BORDER = RGBColor(39, 53, 82)    # #273552
    TEXT_WHITE = RGBColor(248, 250, 252)  # #F8FAFC
    TEXT_MUTED = RGBColor(148, 163, 184)  # #94A3B8
    ACCENT_BLUE = RGBColor(56, 189, 248)  # #38BDF8 Sky Blue
    ACCENT_GREEN = RGBColor(52, 211, 153) # #34D399 Mint Green
    ACCENT_ORANGE = RGBColor(251, 146, 60)# #FB923C Coral/Amber
    ACCENT_PURPLE = RGBColor(192, 132, 252) # #C084FC Lavender

    def set_slide_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, badge_text, title_text, subtitle_text=""):
        # Badge
        badge_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.5), Inches(3.2), Inches(0.4))
        badge_box.fill.solid()
        badge_box.fill.fore_color.rgb = RGBColor(30, 41, 59)
        badge_box.line.color.rgb = ACCENT_BLUE
        badge_box.line.width = Pt(1)
        tf_badge = badge_box.text_frame
        tf_badge.word_wrap = True
        p_badge = tf_badge.paragraphs[0]
        p_badge.text = badge_text.upper()
        p_badge.font.size = Pt(11)
        p_badge.font.bold = True
        p_badge.font.color.rgb = ACCENT_BLUE
        p_badge.alignment = PP_ALIGN.CENTER

        # Title
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.95), Inches(11.7), Inches(0.8))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE

        # Subtitle
        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.size = Pt(13)
            p2.font.color.rgb = TEXT_MUTED

    def add_card(slide, left, top, width, height, title, items, accent_color=ACCENT_BLUE, icon="📌"):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1.2)

        # Content box inside
        content_box = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.2), Inches(width - 0.5), Inches(height - 0.4))
        tf = content_box.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.05)
        tf.margin_bottom = Inches(0.05)
        tf.margin_left = Inches(0.05)
        tf.margin_right = Inches(0.05)

        # Card Title
        p_title = tf.paragraphs[0]
        p_title.text = f"{icon}  {title}"
        p_title.font.size = Pt(16)
        p_title.font.bold = True
        p_title.font.color.rgb = accent_color
        p_title.space_after = Pt(12)

        # Items
        for item in items:
            p_item = tf.add_paragraph()
            p_item.text = f"•  {item}"
            p_item.font.size = Pt(12.5)
            p_item.font.color.rgb = TEXT_WHITE
            p_item.space_after = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 1: PROBLEM STATEMENT
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1)
    add_header(s1, "1️⃣ Problem Statement", "The Crisis in Competitive Exam Tracking", "Millions of aspirants struggle with scattered, unpredictable, and buried notifications.")
    
    add_card(s1, 0.8, 2.0, 3.6, 4.8, "Extreme Fragmentation", [
        "Major national exams (JEE, NEET, GATE, CAT, UPSC, SSC) operate on separate, clunky government portals.",
        "Crucial updates are hidden inside unindexed PDFs with vague titles like 'Public Notice 14-B'.",
        "Students are forced to refresh 8-12 different sites daily or depend on unofficial sources."
    ], ACCENT_ORANGE, "🌪️")

    add_card(s1, 4.8, 2.0, 3.6, 4.8, "Missed Deadlines & Career Loss", [
        "Over 300,000+ candidates miss critical deadlines annually (correction windows, fee payment, admit card downloads).",
        "Application windows can be as short as 48-72 hours for corrections.",
        "Missing a single deadline means losing 1 entire year of rigorous preparation and career momentum."
    ], ACCENT_ORANGE, "⏳")

    add_card(s1, 8.8, 2.0, 3.6, 4.8, "Misinformation & Anxiety", [
        "YouTube 'clickbait' thumbnails and viral WhatsApp rumors trigger panic and false exam date rumors.",
        "No existing platform guarantees zero-hallucination verification from official government servers.",
        "Coaching institutes cannot monitor individual candidate deadlines at scale."
    ], ACCENT_ORANGE, "⚠️")

    # -------------------------------------------------------------
    # SLIDE 2: PROPOSED SOLUTION
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2)
    add_header(s2, "2️⃣ Proposed Solution", "ExamAlert AI — Intelligent Official Exam Guardian", "Real-time official website monitoring with guaranteed zero-hallucination AI extraction.")

    add_card(s2, 0.8, 2.0, 3.6, 4.8, "Official Source Sentinel", [
        "Continuous 24/7 background scrapers dedicated to NTA, IIT GATE, IIM CAT, UPSC, and SSC official portals.",
        "Strict Source-Link Mandate: Every alert links directly to verified government domains (.gov.in, .ac.in, .nic.in).",
        "Zero hallucination guarantee — never fabricates dates or unannounced schedules."
    ], ACCENT_GREEN, "🛡️")

    add_card(s2, 4.8, 2.0, 3.6, 4.8, "SHA-256 Change Detection", [
        "Real-time content-hash engine detects authentic portal updates within seconds of publication.",
        "Eliminates noise, duplicate alerts, and minor cosmetic HTML changes.",
        "Dual AI parser extracts structured deadlines, fees, and category metadata in standardized JSON."
    ], ACCENT_GREEN, "⚡")

    add_card(s2, 8.8, 2.0, 3.6, 4.8, "Multi-Channel & Timeline", [
        "Instant Telegram Bot push notifications & responsive HTML email bulletins directly to the student.",
        "Automated countdown urgency reminders (7-day, 3-day, and 1-day warnings).",
        "Interactive Exam Timeline tracking: Registration → Deadline → Correction → Admit Card → Exam → Results."
    ], ACCENT_GREEN, "📱")

    # -------------------------------------------------------------
    # SLIDE 3: TARGET USERS
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3)
    add_header(s3, "3️⃣ Target Users", "Who Needs ExamAlert AI?", "Serving the 30M+ student examination ecosystem across India.")

    add_card(s3, 0.8, 2.0, 3.6, 4.8, "Competitive Exam Aspirants", [
        "30M+ candidates appearing annually for engineering, medical, civil service, management, and government jobs.",
        "Need peace of mind to focus 100% on studying without the constant dread of missing notifications.",
        "Value instant mobile alerts (Telegram/Email) with direct links to apply."
    ], ACCENT_BLUE, "🎓")

    add_card(s3, 4.8, 2.0, 3.6, 4.8, "Coaching Institutes & Mentors", [
        "Over 100,000+ coaching centres (Kota, Delhi, Hyderabad, and tier-2/3 hubs).",
        "Need automated tracking for batches of 500-10,000 students without manual monitoring staff.",
        "Helps counselors guide students through application corrections and counselling phases."
    ], ACCENT_BLUE, "🏫")

    add_card(s3, 8.8, 2.0, 3.6, 4.8, "Parents & Working Aspirants", [
        "Working professionals preparing for GATE, CAT, and UPSC who have limited daily personal time.",
        "Concerned parents tracking their children's critical college admission deadlines.",
        "Requires simplified, unambiguous deadline summaries on a single dashboard."
    ], ACCENT_BLUE, "👨‍👩‍👧")

    # -------------------------------------------------------------
    # SLIDE 4: TECHNICAL APPROACH
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4)
    add_header(s4, "4️⃣ Technical Approach", "Architecture & Robust Technical Pipeline", "Built with modern fullstack technologies engineered for speed, accuracy, and reliability.")

    add_card(s4, 0.8, 2.0, 3.6, 4.8, "Data & AI Pipeline", [
        "Modular Scrapers: BeautifulSoup4 + HTTPX async fetchers targeting dynamic official notice boards.",
        "Change Detector: SHA-256 hashing on normalized DOM structures prevents false positives.",
        "AI Analysis Engine: Dual architecture pairing LLM reasoning with deterministic regex rule engine fallback."
    ], ACCENT_PURPLE, "🧠")

    add_card(s4, 4.8, 2.0, 3.6, 4.8, "Backend & Automation", [
        "FastAPI (Python 3.14): High-performance async RESTful API with auto-generated OpenAPI documentation.",
        "APScheduler: Background task scheduler running continuous health and deadline monitoring.",
        "SQLAlchemy ORM + SQLite/PostgreSQL: Structured schema with User, Exam, Subscription, and Notice entities."
    ], ACCENT_PURPLE, "⚙️")

    add_card(s4, 8.8, 2.0, 3.6, 4.8, "Frontend & Delivery", [
        "React 19 + Vite + Tailwind CSS: Sleek dark-mode interface with Lucide animated iconography.",
        "Multi-Channel Dispatcher: Telegram Bot API + SMTP/Resend HTML email engine.",
        "Hackathon Demo Simulator: Interactive 5-step live pipeline visualization built into the UI."
    ], ACCENT_PURPLE, "💻")

    # -------------------------------------------------------------
    # SLIDE 5: MARKET & BUSINESS POTENTIAL
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s5)
    add_header(s5, "5️⃣ Market & Business Potential", "Massive Market Size & Monetization Strategy", "Tapping into India's $4.5B+ test preparation and competitive examination market.")

    add_card(s5, 0.8, 2.0, 3.6, 4.8, "Total Addressable Market", [
        "TAM: 30M+ candidates sit for national & state competitive entrance exams every year in India.",
        "SAM: 12M+ digitally active students preparing with online platforms, apps, and Telegram groups.",
        "EdTech test-prep market projected to reach $8.5B by 2030."
    ], ACCENT_GREEN, "📈")

    add_card(s5, 4.8, 2.0, 3.6, 4.8, "B2C Subscription (Freemium)", [
        "Free Tier: Core email alerts for up to 2 exams + web timeline dashboard.",
        "Pro Tier (₹49-₹99/month): Unlimited exams, instant Telegram/WhatsApp push alerts, SMS alerts, and AI document checklist.",
        "Affordable pricing with viral peer-to-peer referral loops among students."
    ], ACCENT_GREEN, "💳")

    add_card(s5, 8.8, 2.0, 3.6, 4.8, "B2B Institute Licensing", [
        "Institutional SaaS Tier (₹10k-₹50k/year per institute) for coaching academies and colleges.",
        "White-labeled student notification portals and bulk broadcast integrations.",
        "Affiliate Partnerships: Verified application preparation materials, test series, and college counselling services."
    ], ACCENT_GREEN, "🤝")

    # -------------------------------------------------------------
    # SLIDE 6: SCALABILITY & FUTURE
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s6)
    add_header(s6, "6️⃣ Scalability & Future", "Scaling from 8 Exams to 500+ Portals", "Engineered for horizontal scalability across regional states and international exams.")

    add_card(s6, 0.8, 2.0, 3.6, 4.8, "Scalable Infrastructure", [
        "Decoupled Microservice Architecture: Distributed scraper nodes orchestrated with Celery & Redis queues.",
        "Headless Browser Pool: Playwright/Selenium cluster for JavaScript-heavy state government portals.",
        "Elastic Cloud Containerization: Dockerized microservices ready for AWS ECS / Kubernetes auto-scaling."
    ], ACCENT_BLUE, "🌐")

    add_card(s6, 4.8, 2.0, 3.6, 4.8, "Geographic & Exam Expansion", [
        "Phase 1: Top 15 Central & National Exams (JEE, NEET, GATE, CAT, UPSC, SSC, CUET, NDA).",
        "Phase 2: State-level exams (MHT-CET, KCET, WBJEE, UPPSC, TS EAMCET) across 28 states.",
        "Phase 3: International higher education & exams (GRE, GMAT, IELTS, TOEFL, SAT, USMLE)."
    ], ACCENT_BLUE, "🚀")

    add_card(s6, 8.8, 2.0, 3.6, 4.8, "Enterprise Partner Ecosystem", [
        "Open API for EdTech leaders (PhysicsWallah, Unacademy, Allen) to embed real-time alert widgets into their apps.",
        "State Government & University Notification Feeds integration.",
        "Automated calendar sync (Google Calendar, Apple Calendar, Outlook) for every active deadline."
    ], ACCENT_BLUE, "🔗")

    # -------------------------------------------------------------
    # SLIDE 7: IF WE HAD MORE TIME
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s7)
    add_header(s7, "7️⃣ If We Had More Time", "What We Would Build Next", "Next-phase features to transform ExamAlert AI into the definitive student companion.")

    add_card(s7, 0.8, 2.0, 3.6, 4.8, "WhatsApp Two-Way AI Bot", [
        "WhatsApp Business Cloud API integration: Students can chat directly ('When is JEE correction window?') and receive instant answers.",
        "Send interactive confirmation buttons and admit card PDF links directly in WhatsApp chat.",
        "Eliminates app-install barrier for tier-3 and rural candidates."
    ], ACCENT_ORANGE, "💬")

    add_card(s7, 4.8, 2.0, 3.6, 4.8, "AI Document & Eligibility Scanner", [
        "Computer Vision & OCR engine to verify candidate documents (caste certificates, photos, signatures) against official NTA/UPSC specifications.",
        "Detects resolution, format, and dimension errors before form upload, eliminating 90% of application rejections.",
        "Automated personalized eligibility checker matching age/qualifications to official criteria."
    ], ACCENT_ORANGE, "📄")

    add_card(s7, 8.8, 2.0, 3.6, 4.8, "Multilingual Regional Support", [
        "Real-time official notice translation into 10 Indian regional languages (Hindi, Tamil, Telugu, Marathi, Bengali, etc.).",
        "Voice-based query assistant for vernacular students.",
        "Automated post-exam counselling guidance & college cutoff prediction tracker based on official rank lists."
    ], ACCENT_ORANGE, "🌍")

    prs.save(filename)
    print(f"Deck saved successfully: {filename}")

if __name__ == "__main__":
    create_deck()
